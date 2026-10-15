"""Recompute Figure 7 / Table 3 on the unchanged 231 workbook CT pairs.

Run from the repository root: python data_process/revise_fig7_table3.py
No fitting, force recalibration, filtering, steady-window changes, or new
condition pairing is performed. Inferred raw-to-bin lineage is used ONLY for
input metadata and audits, never to replace the cached comparison CT pairs.
Density / actual-rpm upstream inconsistencies remain explicitly unresolved.
"""
from __future__ import annotations

import hashlib
import json
import platform
import shutil
import sys
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.lines import Line2D
from openpyxl import load_workbook

sys.dont_write_bytecode = True
import plot_theory_experiment_comparison as legacy

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data_process" / "fig7_table3_revision"
ARCHIVE = OUT / "before_20261013"
ANGLES = legacy.ANGLE_LEVELS
RPMS = legacy.RPM_LEVELS
RAW = ROOT / "data" / "旋转测试数据汇总表_升力标定实验残差.xlsx"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def save_json(path: Path, obj: object) -> None:
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")


def csv(frame: pd.DataFrame, name: str) -> None:
    # Preserve all available floating-point precision; round only display tables.
    frame.to_csv(OUT / name, index=False, encoding="utf-8-sig", float_format="%.17g")


def snapshot() -> None:
    """Freeze the current working-tree version, not the Git HEAD version."""
    if (ARCHIVE / "manifest.json").exists():
        return
    files = [ROOT / "template.tex", RAW, legacy.DEFAULT_WORKBOOK,
             ROOT / "data_process" / "plot_theory_experiment_comparison.py",
             ROOT / "data_process" / "plot_rotor_model_validation.py"]
    files += list((ROOT / "data_process").glob("model_validation*"))
    files += list((ROOT / "Image").glob("fig9_rotor_model_diagnostics.*"))
    manifest = {}
    for source in files:
        if source.is_file():
            rel = source.relative_to(ROOT)
            target = ARCHIVE / rel
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(source, target)
            manifest[rel.as_posix()] = sha256(source)
    save_json(ARCHIVE / "manifest.json", manifest)


def lineage(curves: pd.DataFrame, points: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, dict]:
    """Recover existing bin membership and verify against cached means/counts.

    round(10*J/group_max_J) reproduces ALL 231 cached n, J and both CT means
    within the workbook's six-decimal export rounding. This is source-lineage
    reconstruction, NOT a new aggregation used in the primary statistics.
    """
    points = points.copy()
    points["source_excel_row"] = np.arange(2, len(points) + 2)
    points["bin_id"] = points.groupby(["rpm_nom", "alpha_deg"])["J"].transform(
        lambda z: np.rint(10 * z / z.max()).astype(int))
    points["condition_id"] = [f"{int(r)}_{int(a):02d}_{int(b):02d}" for r, a, b in
                              points[["rpm_nom", "alpha_deg", "bin_id"]].itertuples(index=False, name=None)]
    rows = []
    for (rpm, angle, b), g in points.groupby(["rpm_nom", "alpha_deg", "bin_id"], sort=True):
        rows.append({
            "rpm_nom": rpm, "alpha_deg": angle, "bin_id": b,
            "n_source_records": len(g),
            "rpm_actual_mean": g.rpm_mean.mean(), "rpm_actual_min": g.rpm_mean.min(),
            "rpm_actual_max": g.rpm_mean.max(), "V_actual_mean_m_s": g.v_a.mean(),
            "V_actual_min_m_s": g.v_a.min(), "V_actual_max_m_s": g.v_a.max(),
            "J_reconstructed": g.J.mean(), "CT_exp_reconstructed": g.C_T_rotor.mean(),
            "CT_theory_reconstructed": g.C_T_theory.mean(),
            "CT_theory_existing_rpmcorr_mean": g.C_T_theory_rpmcorr.mean(),
            "source_excel_rows": ";".join(map(str, g.source_excel_row)),
            "source_runs_sweeps": ";".join(f"{r}:{sw}" for r, sw in zip(g.run, g.sweep)),
            "old_mape_excluded_records": int((~g.thrust_ok).sum()),
        })
    meta = pd.DataFrame(rows)
    c = curves.copy()
    c["source_curve_excel_row"] = np.arange(2, len(c) + 2)
    c["bin_id"] = c.groupby(["rpm_nom", "alpha_deg"]).cumcount()
    merged = c.merge(meta, on=["rpm_nom", "alpha_deg", "bin_id"], validate="one_to_one", how="outer", indicator=True)
    if not (merged["_merge"] == "both").all():
        raise ValueError("Existing bin membership could not be recovered; do not invent metadata.")
    merged = merged.drop(columns="_merge").sort_values("source_curve_excel_row").reset_index(drop=True)
    deviations = {col: float((merged[col] - merged[reconstructed]).abs().max()) for col, reconstructed in
                  [("J", "J_reconstructed"), ("C_T_exp", "CT_exp_reconstructed"), ("C_T_theory", "CT_theory_reconstructed")]}
    if any(v > 1.01e-6 for v in deviations.values()) or not (merged.n == merged.n_source_records).all():
        raise ValueError(f"Source lineage failed cached-value reconciliation: {deviations}")
    merged["condition_id"] = [f"{int(r)}_{int(a):02d}_{int(b):02d}" for r, a, b in
                              merged[["rpm_nom", "alpha_deg", "bin_id"]].itertuples(index=False, name=None)]
    return merged, points, {"rule": "rint(10*J/max(J) within rpm_nom,alpha_deg); metadata only",
                            "max_cached_mean_deviation": deviations, "exact_n_reconciliation": True}


def errors(frame: pd.DataFrame, theory_column: str = "C_T_theory") -> pd.DataFrame:
    f = frame.copy()
    f["residual_CT"] = f.C_T_exp - f[theory_column]
    f["abs_residual_CT"] = f.residual_CT.abs()
    f["squared_residual_CT"] = f.residual_CT ** 2
    f["relative_error_usable"] = np.isfinite(f.C_T_exp) & (f.C_T_exp > 0)
    f["signed_relative_error_pct"] = np.nan
    ok = f.relative_error_usable
    f.loc[ok, "signed_relative_error_pct"] = 100 * f.loc[ok, "residual_CT"] / f.loc[ok, "C_T_exp"]
    f["absolute_relative_error_pct"] = f.signed_relative_error_pct.abs()
    group_mean = f.groupby("alpha_deg").C_T_exp.transform("mean")
    f["angle_CT_exp_mean"] = group_mean
    f["overall_CT_exp_mean"] = f.C_T_exp.mean()
    f["angle_normalized_residual_pct"] = 100 * f.residual_CT / group_mean
    f["overall_normalized_residual_pct"] = 100 * f.residual_CT / f.C_T_exp.mean()
    return f


def metrics(frame: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for angle, g in [(a, frame.loc[frame.alpha_deg == a]) for a in ANGLES] + [(None, frame)]:
        mean = float(g.C_T_exp.mean())
        rmse = float(np.sqrt(g.squared_residual_CT.mean()))
        rel = g.signed_relative_error_pct
        complete = bool(g.relative_error_usable.all())
        rows.append({"group": "全部" if angle is None else str(angle), "alpha_deg": angle,
                     "N": len(g), "mean_CT_exp": mean, "RMSE_CT": rmse,
                     "NRMSE_pct": 100 * rmse / mean,
                     "normalized_max_abs_error_pct": 100 * g.abs_residual_CT.max() / mean,
                     "MAPE_pct": float(rel.abs().mean()) if complete else np.nan,
                     "median_signed_relative_error_pct": float(rel.median()) if complete else np.nan,
                     "relative_error_unusable_count": int((~g.relative_error_usable).sum())})
    return pd.DataFrame(rows)


def distributions(frame: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for angle, g in frame.groupby("alpha_deg", sort=True):
        z = g.signed_relative_error_pct
        q1, med, q3 = z.quantile([.25, .5, .75])
        inside = z.between(q1 - 1.5 * (q3-q1), q3 + 1.5 * (q3-q1))
        rows.append({"alpha_deg": angle, "N": len(g), "q1_pct": q1, "median_pct": med,
                     "q3_pct": q3, "whisker_low_pct": z[inside].min(), "whisker_high_pct": z[inside].max(),
                     "min_pct": z.min(), "max_pct": z.max(), "negative_count": int((z<0).sum()),
                     "positive_count": int((z>0).sum()), "zero_count": int((z==0).sum())})
    return pd.DataFrame(rows)


def plot_figure(frame: pd.DataFrame) -> None:
    legacy.configure_matplotlib()
    plt.rcParams["svg.fonttype"] = "path"
    fig = plt.figure(figsize=(3.6, 6.25))
    gs = fig.add_gridspec(2, 1, height_ratios=[1.40, 1.0], hspace=.32)
    a, b = fig.add_subplot(gs[0]), fig.add_subplot(gs[1])
    low = min(frame.C_T_theory.min(), frame.C_T_exp.min())
    high = max(frame.C_T_theory.max(), frame.C_T_exp.max())
    pad = .055 * (high-low)
    lim = (low-pad, high+pad)
    x = np.linspace(*lim, 500)
    a.plot(x, x, color=".20", lw=.9, ls="--", zorder=1)
    # Relative error denominator is y (measurement): NEVER use (1 +/- p)*x.
    for p, color, ls, lw in [(.10, ".48", "-.", .70), (.05, ".72", ":", .60)]:
        for sign in (-1, 1):
            a.plot(x, x/(1+sign*p), color=color, ls=ls, lw=lw, zorder=1)
    for rpm in RPMS:
        g = frame.loc[frame.rpm_nom == rpm]
        a.scatter(g.C_T_theory, g.C_T_exp, s=15, marker=legacy.RPM_MARKERS[rpm],
                  facecolors="none", edgecolors=legacy.RPM_COLORS[rpm], lw=.65, alpha=.8, zorder=3)
    a.set(xlim=lim, ylim=lim)
    a.set_aspect("equal", adjustable="box")
    a.set_xlabel(r"理论 $C_T$ 值", fontfamily=legacy.CJK_FONT_FAMILY)
    a.set_ylabel(r"实测 $C_T$ 值", fontfamily=legacy.CJK_FONT_FAMILY)
    a.set_title("(a)", loc="left", fontweight="bold")
    a.grid(color=".90", lw=.4, ls=":")
    handles = [Line2D([], [], marker=legacy.RPM_MARKERS[r], markerfacecolor="none",
                      color=legacy.RPM_COLORS[r], lw=0, markersize=4.5, label=f"{r} rpm") for r in RPMS]
    a.legend(handles=handles, loc="upper left", frameon=False, handletextpad=.4, fontsize=7)
    guides = [Line2D([], [], color=col, ls=ls, lw=lw, label=label) for col, ls, lw, label in
              [(".20", "--", .9, "1:1"), (".48", "-.", .7, r"$\pm10\%$"), (".72", ":", .6, r"$\pm5\%$")]]
    a.add_artist(a.get_legend())
    a.legend(handles=guides, loc="lower right", frameon=False, handlelength=2.4, fontsize=7,
             bbox_to_anchor=(1.0, .15))
    count = int((frame.absolute_relative_error_pct <= 10).sum())
    a.text(.98, .025, f"$N={len(frame)}$\n" + rf"$|e|\leq10\%$: {100*count/len(frame):.1f}% ({count}/{len(frame)})",
           transform=a.transAxes, ha="right", va="bottom", fontsize=7,
           bbox={"facecolor": "white", "edgecolor": "none", "alpha": .9, "pad": 1.5})
    values = [frame.loc[frame.alpha_deg == angle, "signed_relative_error_pct"].to_numpy() for angle in ANGLES]
    b.axhline(0, color=".25", lw=.85, ls="--", zorder=1)
    for threshold in (-10, 10):
        b.axhline(threshold, color=".65", lw=.65, ls="-.", zorder=1)
    for threshold in (-5, 5):
        b.axhline(threshold, color=".80", lw=.55, ls=":", zorder=1)
    b.boxplot(values, positions=np.arange(7), widths=.65, patch_artist=True, whis=1.5, showfliers=False,
              boxprops={"facecolor": "#DFEAF0", "edgecolor": ".3", "linewidth": .7},
              medianprops={"color": ".1", "linewidth": 1.3},
              whiskerprops={"color": ".3", "linewidth": .7}, capprops={"color": ".3", "linewidth": .7})
    for pos, angle in enumerate(ANGLES):
        for rpm, offset in zip(RPMS, (-.19, 0, .19)):
            g = frame.loc[(frame.alpha_deg==angle)&(frame.rpm_nom==rpm)].sort_values("J")
            # Small deterministic offsets spread coincident points without changing y.
            jitter = np.linspace(-.045, .045, len(g))
            b.scatter(pos + offset + jitter, g.signed_relative_error_pct, s=10,
                      marker=legacy.RPM_MARKERS[rpm], facecolors="none", edgecolors=legacy.RPM_COLORS[rpm],
                      lw=.55, alpha=.85, zorder=3)
    extent = max(10, frame.absolute_relative_error_pct.max())
    b.set_ylim(-1.17*extent, 1.17*extent)
    b.set_xticks(np.arange(7), labels=ANGLES)
    b.set_xlabel(r"桨盘攻角 $\alpha_p$ ($^\circ$)", fontfamily=legacy.CJK_FONT_FAMILY)
    b.set_ylabel("有符号相对误差 (%)", fontfamily=legacy.CJK_FONT_FAMILY)
    b.set_title("(b)", loc="left", fontweight="bold")
    b.text(.98, .96, "33 conditions per angle", transform=b.transAxes, ha="right", va="top", fontsize=7)
    b.grid(axis="y", color=".90", lw=.4, ls=":", zorder=0)
    fig.subplots_adjust(left=.19, right=.97, top=.965, bottom=.08)
    for suffix in ("pdf", "svg", "png"):
        dest = OUT / f"fig7_revised.{suffix}"
        fig.savefig(dest, dpi=600, metadata={"Creator": "revise_fig7_table3.py"} if suffix == "pdf" else None)
    plt.close(fig)
    for suffix in ("pdf", "png"):
        shutil.copy2(OUT / f"fig7_revised.{suffix}", ROOT / "Image" / f"fig9_rotor_model_diagnostics.{suffix}")


def raw_audit() -> dict:
    w = load_workbook(RAW, data_only=True, read_only=True)
    density_rows = []
    for name in w.sheetnames[:3]:
        rows = list(w[name].values)
        head = rows[1]
        start = head.index("rpm")
        for idx, row in enumerate(rows[2:], start=3):
            if len(row) <= start+6 or not isinstance(row[start], (float,int)):
                continue
            rpm, omega, angle, Omega, force, raw_ct, ct = row[start:start+7]
            if all(isinstance(x, (int,float)) and x>0 for x in [Omega,force,ct]):
                rho = force / (ct * np.pi * (.4572/2)**4 * Omega**2)
                density_rows.append({"sheet": name, "excel_row": idx, "rpm": rpm, "model_angle_field": angle,
                                     "T_N": force, "CT_rotor": ct, "Omega_rad_s": Omega,
                                     "rho_implied_kg_m3": rho})
    cal = w[w.sheetnames[3]]
    constants = {str(k): v for k,v,*_ in list(cal.values)[:8] if isinstance(v,(float,int))}
    w.close()
    csv(pd.DataFrame(density_rows), "theory_density_audit.csv")
    return {"experiment_calibration_constants": constants,
            "theory_rho_implied_min": min(r["rho_implied_kg_m3"] for r in density_rows),
            "theory_rho_implied_max": max(r["rho_implied_kg_m3"] for r in density_rows),
            "theory_rows_checked": len(density_rows),
            "diameter_m": .4572, "radius_m": .2286,
            "theory_radius_assumption": "rho back-calculation assumes the same 0.2286 m rotor radius; upstream theory setup script is unavailable",
            "status": "UNRESOLVED upstream density mismatch; no rescaling/model rerun performed"}


def markdown_table(frame: pd.DataFrame, columns: list[tuple[str,str]]) -> str:
    lines = ["| " + " | ".join(label for key,label in columns) + " |",
             "| " + " | ".join("---" for _ in columns) + " |"]
    for record in frame.to_dict("records"):
        values = []
        for key,label in columns:
            v=record[key]
            values.append(str(v) if isinstance(v,str) else str(int(v)) if key in {"N","old_N"} else f"{v:.2f}")
        lines.append("| " + " | ".join(values) + " |")
    return "\n".join(lines)


def write_table(new: pd.DataFrame, old: pd.DataFrame) -> pd.DataFrame:
    compare = new[["group","N","NRMSE_pct","normalized_max_abs_error_pct","MAPE_pct"]].copy()
    for target,source in [("old_N","n"),("old_NRMSE_pct","rmse_norm_pct"),
                          ("old_normalized_max_abs_error_pct","maxerr_norm_pct"),("old_MAPE_pct","mape_pct")]:
        compare[target] = old[source].to_numpy()
    for col in ["NRMSE_pct","normalized_max_abs_error_pct","MAPE_pct"]:
        compare[f"delta_{col}_percentage_points"] = compare[col] - compare[f"old_{col}"]
    csv(compare,"table3_old_new_comparison.csv")
    rows = [f"{r.group} & {r.NRMSE_pct:.2f} & {r.normalized_max_abs_error_pct:.2f} & {r.MAPE_pct:.2f} \\\\" for r in new.itertuples()]
    tex = "\n".join([
        r"\begin{center}", r"\vbox{\centering{\small 表3\quad 不同桨盘攻角下的旋翼升力模型误差统计",
        r"\\ Table\,3 \quad Error statistics of the rotor lift model at different disk angles of attack}\vskip2mm",
        r"{\scriptsize\centering", r"\setlength{\tabcolsep}{3.5pt}", r"\begin{tabular}{@{}crrr@{}}", r"\toprule",
        r"$\alpha_p/(^\circ)$ & NRMSE/\% & MaxErr$_{\mathrm{norm}}$/\% & MAPE/\% \\", r"\midrule", *rows,
        r"\bottomrule",r"\end{tabular}}",r"\vskip1mm",
        r"\parbox{0.98\columnwidth}{\scriptsize 注：每个攻角33个工况，全部231个工况，各工况等权。NRMSE与归一化最大绝对误差的分母均为对应统计组实测$C_T$均值；MAPE先以各工况实测$C_T$归一化再取算术平均。保留现有$C_T$配对；上游密度与转速修正一致性尚待核实。}",
        r"}",r"\end{center}",""])
    (OUT/"table3_revised.tex").write_text(tex,encoding="utf-8")
    shutil.copy2(OUT/"table3_revised.tex", ROOT/"data_process"/"model_validation_table.tex")
    return compare


def manuscript_spans(data: bytes) -> list[tuple[int,int]]:
    spans=[]
    for anchor in [b"{Image/fig9_rotor_model_diagnostics.pdf}", "{\\small 表3\\quad".encode('utf-8')]:
        if data.count(anchor)!=1:
            raise ValueError("Expected a unique Figure 7 / Table 3 source block; no broad text replacement.")
        pos=data.index(anchor)
        start=data.rfind(b"\\begin{center}",0,pos)
        end=data.index(b"\\end{center}",pos)+len(b"\\end{center}")
        spans.append((start,end))
    return sorted(spans)


def outside_manuscript_spans(data: bytes) -> bytes:
    spans=manuscript_spans(data)
    return data[:spans[0][0]]+b"<FIG7>"+data[spans[0][1]:spans[1][0]]+b"<TABLE3>"+data[spans[1][1]:]


def update_manuscript_artifacts() -> None:
    """Only the Figure 7 caption block and inline Table 3, never prose."""
    path=ROOT/'template.tex'
    original=path.read_bytes()
    nl=b'\r\n' if b'\r\n' in original else b'\n'
    caption="\n".join([
        r"\begin{center}",r"\vbox{\centering",
        r"\includegraphics[width=7.5cm]{Image/fig9_rotor_model_diagnostics.pdf}\\[1mm]",
        r"{\small 图7\quad 升力系数一致性与攻角分区相对误差分布",
        r"\\ Fig.\,7 \quad Thrust coefficient agreement and relative-error distributions by disk angle of attack}\\[1mm]",
        r"\parbox{0.98\columnwidth}{\scriptsize 注：(a)参考线为1:1及相对实测值的$\pm10\%$、$\pm5\%$误差边界。(b)有符号相对误差以实测$C_T$为分母，正值表示模型低估，负值表示模型高估。每个攻角合并三个转速档位的33个不同工况；箱体为第一至第三四分位数，横线为中位数，须线延伸至1.5倍四分位距内的最远观测值，叠加显示全部工况。箱体表示工况间分布，不表示均值置信区间或33次重复试验。}",
        r"}",r"\end{center}"])
    replacement=[caption.encode('utf-8').replace(b'\n',nl),
                 (OUT/'table3_revised.tex').read_bytes().rstrip(b'\r\n').replace(b'\r\n',b'\n').replace(b'\n',nl)]
    updated=original
    for (start,end),new in reversed(list(zip(manuscript_spans(original),replacement))):
        updated=updated[:start]+new+updated[end:]
    if outside_manuscript_spans(original)!=outside_manuscript_spans(updated):
        raise AssertionError("Unexpected manuscript edit outside Figure 7 / Table 3.")
    path.write_bytes(updated)
    save_json(OUT/'manuscript_scope_verification.json',{
        "allowed_edits":["Figure 7 caption block","Table 3 inline block"],
        "outside_blocks_sha256":hashlib.sha256(outside_manuscript_spans(updated)).hexdigest(),
        "current_source_sha256":sha256(path),
        "outside_blocks_byte_identical":True,
        "paper_recompiled":False})


def main() -> None:
    workbook = legacy.DEFAULT_WORKBOOK
    OUT.mkdir(parents=True, exist_ok=True)
    snapshot()
    protected = [RAW, workbook, ROOT/"data_process"/"plot_theory_experiment_comparison.py",
                 ROOT/"data_process"/"plot_rotor_model_validation.py"]
    protected_hashes = {str(p.relative_to(ROOT)):sha256(p) for p in protected}
    s = legacy.load_sheets(workbook)
    c,p,lin = lineage(s["curves"],s["points"])
    if not np.isfinite(c[["C_T_exp","C_T_theory"]].to_numpy()).all():
        raise ValueError("Nonfinite CT pairs: retain and investigate; no automatic exclusion.")
    if c.duplicated(["rpm_nom","alpha_deg","J"]).any():
        raise ValueError("Duplicate comparison conditions; no silent deduplication.")
    if set(c.rpm_nom)!=set(RPMS) or set(c.alpha_deg)!=set(ANGLES):
        raise ValueError("Unexpected operating envelope; no padding of missing conditions.")
    f = errors(c)
    new = metrics(f)
    old = legacy.build_angle_metrics(s["points"])
    old_rpm = legacy.build_point_metrics(s["points"])
    csv(old,"table3_legacy_reproduced.csv")
    csv(old_rpm,"legacy_metrics_by_rpm.csv")
    csv(new,"table3_revised_unrounded.csv")
    csv(distributions(f),"relative_error_distributions_by_angle.csv")
    compare = write_table(new,old)
    update_manuscript_artifacts()
    # Per-condition inputs, CT pairs, provenance, and all error definitions.
    csv(f,"condition_details.csv")
    csv(p,"source_record_lineage.csv")
    csv(c.loc[c.n!=4],"conditions_with_nonfour_source_records.csv")
    csv(p.loc[~p.thrust_ok],"legacy_mape_excluded_records.csv")
    csv(f.loc[~f.relative_error_usable],"relative_error_unusable_conditions.csv")
    signal_review = f[["condition_id", "rpm_nom", "alpha_deg", "J", "V_actual_mean_m_s",
                       "C_T_exp", "C_T_exp_std"]].copy()
    signal_review["CT_to_branch_spread_ratio_NOT_sensor_SNR"] = signal_review.C_T_exp / signal_review.C_T_exp_std
    csv(signal_review.sort_values("CT_to_branch_spread_ratio_NOT_sensor_SNR"), "signal_spread_review_NOT_noise_test.csv")
    # Existing alternate CT column is sensitivity only; never switches primary CT.
    sensitivity = metrics(errors(c,"CT_theory_existing_rpmcorr_mean"))
    csv(sensitivity,"sensitivity_existing_rpmcorr_NOT_primary.csv")
    raw = raw_audit()
    audit = {"source_workbook":str(workbook.relative_to(ROOT)), "source_sheet":"分档曲线_可直接作图",
             "N_expected":231,"N_actual":len(c),"counts_by_angle":{str(a):int((c.alpha_deg==a).sum()) for a in ANGLES},
             "counts_by_rpm_angle":[{"rpm":int(r),"angle":int(a),"N":len(g)} for (r,a),g in c.groupby(["rpm_nom","alpha_deg"])],
             "duplicate_condition_keys":int(c.duplicated(["rpm_nom","alpha_deg","J"]).sum()),
             "duplicate_source_records":int(s['points'].duplicated().sum()),
             "missing_condition_keys":[[r,a,b] for r in RPMS for a in ANGLES for b in range(11) if not ((c.rpm_nom==r)&(c.alpha_deg==a)&(c.bin_id==b)).any()],
             "N_source_records":len(p),"source_count_per_bin":{str(k):int(v) for k,v in c.n.value_counts().sort_index().items()},
             "source_sequence_shortfalls":[{"rpm_nom":int(r),"alpha_deg":int(a),"run":int(run),"sweep":sw,"available":len(g),"expected_if_full":11 if sw=='up' else 10} for (r,a,run,sw),g in p.groupby(['rpm_nom','alpha_deg','run','sweep']) if len(g)!=(11 if sw=='up' else 10)],
             "lineage":lin,"n_nonpositive_CT":int((c.C_T_exp<=0).sum()),"min_CT_exp":float(c.C_T_exp.min()),
             "min_CT_condition":f.loc[f.C_T_exp.idxmin(),["condition_id","rpm_nom","alpha_deg","J","C_T_exp","C_T_exp_std"]].to_dict(),
             "noise_assessment":"NOT VERIFIED: sensor noise/resolution absent. Branch spread is not an independent noise estimate.",
             "raw_dimensionless_audit":raw,
             "CT_convention":"rotor tip-speed CT=T/(rho*pi*R^2*(Omega*R)^2); prop CT multiplier 4/pi^3, not interchangeable",
             "actual_rpm_J_max_abs_reconstruction_deviation":float((p.J-p.v_a/(p.rpm_mean/60*.4572)).abs().max()),
             "prop_to_rotor_CT_max_abs_deviation":float((p.C_T_rotor-p.C_T_prop_u*4/np.pi**3).abs().max()),
             "primary_CT_pairs":"unchanged six-decimal workbook C_T_exp / C_T_theory. Existing rpmcorr column not substituted.",
             "legacy_mape_denominator_N":int(p.thrust_ok.sum()),"legacy_mape_excluded_N":int((~p.thrust_ok).sum()),
             "legacy_extrapolated_flags_true_N":int(p.extrapolated.sum()),
             "environment":{"python":platform.python_version(),"numpy":np.__version__,"pandas":pd.__version__,"matplotlib":matplotlib.__version__}}
    save_json(OUT/"audit.json",audit)
    result = new.iloc[-1]
    special = f.loc[f.absolute_relative_error_pct.idxmax()]
    max_abs = f.loc[f.abs_residual_CT.idxmax()]
    within = {str(th):{"count":int((f.absolute_relative_error_pct<=th).sum()),"N":len(f),
                       "percentage":float(100*(f.absolute_relative_error_pct<=th).sum()/len(f))} for th in [5,10]}
    summary={"primary_basis":"unchanged 231 cached CT pairs; upstream consistency unresolved",
             "overall":result.drop(labels=["alpha_deg"]).to_dict(),"within_abs_relative_error":within,
             "max_absolute_relative_error_condition":special.to_dict(),"max_absolute_CT_error_condition":max_abs.to_dict(),
             "signed_medians_by_angle":{str(a):float(f.loc[f.alpha_deg==a,'signed_relative_error_pct'].median()) for a in ANGLES}}
    save_json(OUT/"summary.json",summary)
    plot_figure(f)
    write_notes(new,old,compare,f,audit,sensitivity,within,special,max_abs)
    for path,hash_before in protected_hashes.items():
        if sha256(ROOT/path)!=hash_before:
            raise AssertionError(f"Protected input changed: {path}")
    save_json(OUT/"protected_inputs_sha256.json",protected_hashes)
    print(new[["group","N","NRMSE_pct","normalized_max_abs_error_pct","MAPE_pct"]].to_string(index=False))
    print(json.dumps(within,ensure_ascii=False))
    print(f"Outputs: {OUT}")


def write_notes(new,old,compare,f,audit,sensitivity,within,special,max_abs):
    table = markdown_table(new,[("group","攻角/°"),("N","N"),("NRMSE_pct","NRMSE/%"),
                               ("normalized_max_abs_error_pct","归一化最大绝对误差/%"),("MAPE_pct","MAPE/%")])
    comparison = markdown_table(compare,[("group","攻角/°"),("old_N","旧N"),("N","新N"),
        ("old_NRMSE_pct","旧NRMSE/%"),("NRMSE_pct","新NRMSE/%"),
        ("old_normalized_max_abs_error_pct","旧最大/%"),("normalized_max_abs_error_pct","新最大/%"),
        ("old_MAPE_pct","旧MAPE/%"),("MAPE_pct","新MAPE/%")])
    (OUT/"table3_revised.md").write_text(table+"\n",encoding="utf-8")
    (OUT/"table3_old_new_comparison.md").write_text(comparison+"\n",encoding="utf-8")
    mid=new.loc[new.alpha_deg.between(15,75)]
    lo,hi=mid.loc[mid.NRMSE_pct.idxmin()],mid.loc[mid.NRMSE_pct.idxmax()]
    end0,end90=new.loc[new.alpha_deg==0].iloc[0],new.loc[new.alpha_deg==90].iloc[0]
    medians="；".join(f"{a}°：{f.loc[f.alpha_deg==a,'signed_relative_error_pct'].median():+.2f}%" for a in ANGLES)
    overall=new.iloc[-1]
    summary = f"""本结果基于工作簿中冻结的231组CT配对；上游空气密度、实际转速理论修正尚未闭合核实，不能宣称已满足同密度、同实际转速的完整一致性要求。

- 实际工况数231，每个攻角33；无缺失工况、无重复工况。底层处理后记录877条，各工况汇总2—5条记录，但新统计每工况等权。
- 绝对相对误差≤5%：{within['5']['count']}/231（{within['5']['percentage']:.2f}%）；≤10%：{within['10']['count']}/231（{within['10']['percentage']:.2f}%）。
- 总体RMSE={overall.RMSE_CT:.10f}；总体NRMSE={overall.NRMSE_pct:.2f}%；MAPE={overall.MAPE_pct:.2f}%；归一化最大绝对误差={overall.normalized_max_abs_error_pct:.2f}%。
- 15°—75°各组NRMSE最小为{lo.NRMSE_pct:.2f}%（{int(lo.alpha_deg)}°），最大为{hi.NRMSE_pct:.2f}%（{int(hi.alpha_deg)}°）。0°为{end0.NRMSE_pct:.2f}%，90°为{end90.NRMSE_pct:.2f}%，均高于全部15°—75°组。
- 各攻角有符号相对误差中位数：{medians}。
- 最大绝对相对误差：{int(special.rpm_nom)} rpm档，实际转速均值{special.rpm_actual_mean:.2f} rpm，{int(special.alpha_deg)}°，来流速度均值{special.V_actual_mean_m_s:.6f} m/s，J={special.J:.6f}；理论CT={special.C_T_theory:.6f}，实测CT={special.C_T_exp:.6f}，残差={special.residual_CT:+.6f}，有符号相对误差={special.signed_relative_error_pct:+.2f}%。
- 最大绝对CT残差是另一工况：{int(max_abs.rpm_nom)} rpm档，实际均值{max_abs.rpm_actual_mean:.2f} rpm，{int(max_abs.alpha_deg)}°，来流速度均值{max_abs.V_actual_mean_m_s:.6f} m/s，J={max_abs.J:.6f}；理论CT={max_abs.C_T_theory:.6f}，实测CT={max_abs.C_T_exp:.6f}，残差={max_abs.residual_CT:+.6f}，相对误差={max_abs.signed_relative_error_pct:+.2f}%。
- 按组中位数方向，支持0°、45°总体偏低估，75°、90°总体偏高估；不意味着这些组中每个工况均为同一偏差方向。不据此归因于失速、变形或尾迹。
"""
    (OUT/"chat_summary_zh.md").write_text(summary,encoding="utf-8")
    paragraph="本统计以每个分档工况的实测与理论升力系数差值为基础。NRMSE的分母为对应统计组的实测升力系数算术均值；归一化最大绝对误差使用同一个分母。MAPE先用各工况的实测升力系数作为分母计算绝对相对误差，再取算术平均。分攻角统计各包含33个工况，总体统计直接使用全部231个工况，各工况等权，不按底层记录数量加权，也不对分组NRMSE作简单平均。有符号相对误差为实测减理论后除以该工况实测值，正值表示模型低估，负值表示模型高估。"
    caption="图7 旋翼升力模型在231个分档工况点上的误差。（a）理论CT与实测CT一致性散点图，颜色及点形区分转速档位；虚线为1∶1参考线，点划线及淡色点线分别为相对实测值的±10%及±5%误差边界。（b）各桨盘攻角下有符号相对误差箱线图，正值表示模型低估，负值表示模型高估。每个攻角合并三个转速档位、各11个来流档位，共33个不同工况；箱体为第一至第三四分位数，横线为中位数，须线延伸至1.5倍四分位距范围内的最远观测值，叠加散点包含全部工况及离群值。箱体反映工况间分布，不表示均值置信区间，也不代表33次重复试验。"
    (OUT/"paper_text_and_caption_zh.md").write_text("适用性说明：当前数字保留原工作簿CT配对；必须先解决audit_report_zh.md中的上游一致性问题，才能将其表述为已核实的统一无量纲化结果。\n\n"+paragraph+"\n\n"+caption+"\n",encoding="utf-8")
    report=f"""# 图7及表3统计核查记录

## Material Passport
模式：学术统计复现与限定范围修订。来源：本项目两个工作簿、现有脚本、template.tex；未向外部服务上传数据。状态：231点冻结配对的计算与绘图已复现；上游密度/转速一致性及噪声量级待核实。

## 范围与定位
README.md仅介绍LaTeX模板与XeLaTeX。template.tex第4.1节使用Image/fig9_rotor_model_diagnostics.pdf，虽文件名含fig9，当前论文编号为图7。表3在template.tex内联写入，并未input生成的data_process/model_validation_table.tex，因此本次同时更新了内联表3块和图7图注块；这两块之外逐字节校验保持不变。主绘图/统计入口是data_process/plot_theory_experiment_comparison.py。另一个plot_rotor_model_validation.py使用较早的残差工作簿，不能作为当前图7/表3的数据替代入口。

- 原始来源：data/旋转测试数据汇总表_升力标定实验残差.xlsx的3000rpm、4000rpm、5000rpm工作表，包含转台角速度、实际旋翼转速、传感器读数、测量均值以及理论力和CT；其升力标定计算、残差明细、残差汇总表保存公式链。
- 当前主来源：data/理论-实验对比结果.xlsx；分档曲线_可直接作图含231点，逐点明细含877条处理后记录，误差汇总_逐工况含21个转速/攻角组合。当前工作簿是数值导出，不含生成公式；未在项目中找到其上游生成脚本，滤波/稳态窗口的更早来源无法复核。本次未重建或改动这些处理步骤。
- before_20261013保存本轮开始时的工作树版本（包含用户未提交修改），并附SHA256。未使用Git版本覆盖用户文件。

## 数量、缺失与记录权重
21个转速/攻角组合各有11个分档点，共231点，每攻角33点。无缺失工况、无重复(rpm,angle,J)键；877条逐点记录无完全重复。分档点底层n分布：204点n=4，22点n=2，4点n=3，1点n=5，总和877；详见conditions_with_nonfour_source_records.csv。最大来流转折点通常只有两次试验的两条记录，没有人为复制成四条。

以每次上行11、下行10条为完整序列参照，应有882条，当前处理后记录少5条：4000rpm/75°/run1下行9条；5000rpm/45°/run1和run2下行各9条；5000rpm/75°/run2下行9条；5000rpm/90°/run1下行9条。这里指处理后序列短缺，不能据此断言原始仪器未采到数据；缺少上游导出脚本，未推测剔除原因、未补点。

source_record_lineage.csv逐条保留已有记录和工作簿行号。依据各组J最大值重建已有分档归属，231组n完全一致，J和CT均值最大偏差均不超过7.5e-7（工作簿六位小数导出精度）。此归属仅用于补充实际转速/来流速度元数据与审计，不用重算或替换主统计的CT配对。各工况J与来流速度是实际均值，11档不意味着各档实际风速严格相同。

## 旧指标的实际来源
旧脚本_pooled_point_stats直接读取逐点明细中的已缓存误差字段；不是对图7的231点统计，也不是简单平均7组指标。

- 旧RMSE_norm：对877条err_norm_pct平方取均值再开方；每条err_norm_pct是CT残差除以所属转速档理论悬停CT再乘100。不同转速分母不同，不能把总体理解为除以一个共同悬停CT。
- 旧MaxErr_norm：上述逐记录err_norm_pct的最大绝对值。
- 旧MAPE：只对thrust_ok=True的869条err_rel_pct取绝对值均值，分母为每条匹配理论CT。其余8条仅从MAPE排除，仍参与旧RMSE与最大误差。
- 旧图7相对残差：231点的(实测CT−理论CT)/理论CT；图7(b)实际画的是绝对CT残差（乘1000显示）。较早残差工作簿也以理论CT为相对残差分母，并在理论CT小于1e-12时返回空值；新脚本没有继承这个epsilon逻辑。
- 分组：按alpha_deg合并三个转速；旧N分别为126、126、126、124、126、124、125，总体877。底层记录等权，记录较多的工况对旧总体贡献较多。

从原始缓存字段直接复现总体：NRMSE={old.iloc[-1].rmse_norm_pct:.12f}%、MAPE={old.iloc[-1].mape_pct:.12f}%、最大归一化误差={old.iloc[-1].maxerr_norm_pct:.12f}%，舍入后分别5.76%、5.08%、19.75%。最大旧误差出现在3000rpm档、90°、实际转速3178.472464rpm、来流9.190950m/s的上行run1记录，err_norm_pct=-19.749097%。这不是新231点最大相对误差工况。

## 确定的说明/计算层级问题及修正记录
1. 旧表与图使用不同统计层级。已按用户要求统一为231个冻结工况点，每点等权；这是统计对象与权重调整，不是为了降误差筛除数据。
2. 旧相对残差及MAPE分母是理论CT。已统一换成实测CT；NRMSE与最大误差换成各统计组实测CT均值。无反推分母，无epsilon，无删除工况。
3. 旧脚本注释把thrust_ok=False等同于超出理论J范围，但8条记录的extrapolated均为False；不能用该注释证明外推。这一说明错误已记录，未改源标记。旧MAPE若取消该筛选为{audit_legacy_all_mape():.8f}%，而原筛选结果为5.08397772%。详见legacy_mape_excluded_records.csv。
4. 旧脚本头部声称使用rpm-corrected理论插值，但实际图表数据复核匹配C_T_theory未修正列；存在已有C_T_theory_rpmcorr列未采用的来源矛盾。因生成算法/意图缺少上游代码，未擅自将其判成可替换的无量纲化修正；主版本保留当前配对。

## CT定义与尚未解决的上游一致性
论文定义是基于Omega和R的旋翼CT。逐点明细C_T_rotor与C_T_prop_u乘4/pi^3一致到六位小数导出精度，最大绝对差{audit['prop_to_rotor_CT_max_abs_deviation']:.3g}；没有将prop系数直接用于新统计。原始旧公式还存在3.14与pi常数近似差异，保留原公式，未二次换算已经转换的CT。

实测J可由v_a/(rpm_mean/60*D)复核，最大偏差{audit['actual_rpm_J_max_abs_reconstruction_deviation']:.3g}；D=0.4572m，R=0.2286m。实测归一化采用各条实际转速，不用档位名代替实际rpm。分档CT为已有系数的平均值，不能用分档平均转速再归一化平均力替代。

原始标定工作表空气密度为1.205kg/m³。假设理论使用同一R=0.2286m，理论T_N/CT/Omega/R反算的{audit['raw_dimensionless_audit']['theory_rows_checked']}条理论记录密度为1.225kg/m³，两者相差约1.66%；理论计算配置未单独保存，不能独立确认其半径参数。仅观察系数不能判断上游是否已对理论力进行了密度转换；当前工作簿没有可追溯的理论力列（Thrust(N)全空）或生成脚本，因此不能确认同密度/同尺寸。也不能仅把理论CT乘1.225/1.205来“修好”：这会隐含固定理论力且改归一化约定，是否合适取决于模型密度和雷诺数处理。本次没有改变模型参数或原始系数。

已有转速修正列的同分档均值作为敏感性计算（不是主表）：总体NRMSE={sensitivity.iloc[-1].NRMSE_pct:.8f}%、MAPE={sensitivity.iloc[-1].MAPE_pct:.8f}%、最大归一化误差={sensitivity.iloc[-1].normalized_max_abs_error_pct:.8f}%。文件sensitivity_existing_rpmcorr_NOT_primary.csv仅量化这一来源差异，不证明完成密度统一，也不用于替换主图表。

结论：新统计计算与冻结231点可复现；“理论与实测同密度、同实际转速”的完整上游一致性不能据当前证据确认。需要作者补充生成对比工作簿的代码及理论密度/转速修正依据后闭合核查。

## 正值与近零检查
231个实测CT均有限且为正，最小0.002383，位于3000rpm档、15°、J=0.407323；对应工况间分支标准差0.000218，二者比约10.93。标准差是分支/重复离散，不能当独立传感器噪声。无确定的非正或无效分母；无epsilon、无隐藏删点。项目没有传感器分辨率/噪声CT，因而不能严格判定是否存在接近噪声量级的近零值；受影响工况集合目前无法判定，而不是默认空集。relative_error_unusable_conditions.csv为空只表示已知非正/非有限问题为零。signal_spread_review_NOT_noise_test.csv列出全部231点的CT与分支标准差比，供人工复核，不将其解释为传感器信噪比或据此剔除点。

## 新表与新旧对照
{table}

{comparison}

差异同时来自：统计对象877记录→231工况、底层记录权重→工况等权、悬停理论系数→组实测CT均值、逐点理论分母→实测分母，以及旧MAPE筛选869记录→全部231工况。CT配对完全保持不变，因此没有新测量筛选、拟合或模型输出修改。新最大归一化误差22.18%与最大逐点绝对相对误差19.56%是不同指标，分母不同。可用对照文件的未舍入差值列核查，不从展示舍入值倒算。

## 图形及正文使用
图7(a)保持等比例同范围，误差边界按实测为分母，±10%与更淡±5%线覆盖数据范围；删除回归式/R²/CCC，无Pearson标注。图7(b)包含全部231点，采用完整对称范围，保留离群点及零线、±10%、±5%线。字体、三档颜色/点形沿用旧脚本。新图在Image同名文件及本目录均有副本；PDF和SVG为矢量，PNG为600dpi。

正文建议文字与完整图注见paper_text_and_caption_zh.md；可直接粘贴聊天的数字见chat_summary_zh.md。template.tex仅更新图7图注块与内联表3块，没有改相邻正文、第4.2节、结论或其他图表；两块之外内容与修改前快照逐字节一致。作者需要随后同步4.1节正文中的旧相对残差分母、回归/相关性文字、旧固定绝对阈值比例、旧残差分布数值、旧表3统计数字以及“各攻角RMSE小于10%、最大误差小于20%”等比较性文字；结论中5.76%、2.46%—5.02%若决定采纳本版也需自行同步。这些正文/结论中的数字没有自动替换。未重编译整篇论文，现有template.pdf仍是本轮之前版本。

## 复现与校验
在项目根目录运行：
```powershell
python data_process/revise_fig7_table3.py
python data_process/verify_fig7_table3_revision.py
```
依赖numpy、pandas、matplotlib、openpyxl；本次精确版本见audit.json。可在任意当前目录用脚本绝对路径运行。脚本只覆盖图7及新表3和专属输出目录，首次运行保存旧版；不执行旧脚本的main，不改其他图形。工作簿本身六位小数导出已经限制输入精度，本轮计算不再提前舍入。

核查涵盖11类统计误读：分组与总体方向已分别报告；不从组中位数推及每个工况；仅描述已测包线；无协变量控制/碰撞变量分析；比例明确分子分母；无前后极端组回归均值推论；底层短缺与旧筛选已披露；无显著性筛选或多重假设检验；固定用户指定分析口径并保留敏感性；无机制因果推断；无反向因果断言。这里没有p值/置信区间/显著性声明。
"""
    (OUT/"audit_report_zh.md").write_text(report,encoding="utf-8")


def audit_legacy_all_mape():
    return float(legacy.load_sheets(legacy.DEFAULT_WORKBOOK)["points"].err_rel_pct.abs().mean())


if __name__ == "__main__":
    main()
