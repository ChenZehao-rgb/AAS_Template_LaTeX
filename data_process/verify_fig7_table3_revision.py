"""Independent, source-bound verification using stdlib math and openpyxl.

Does not import the production metrics/plotting helpers. Recomputes all rows
from workbook bytes, checks unmodified CT pairs and the archived old outputs.
"""
from __future__ import annotations

import csv
import hashlib
import json
import math
import statistics
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data_process" / "fig7_table3_revision"


def read_csv(name):
    with (OUT / name).open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def close(actual, expected, label, tolerance=2e-12):
    if not math.isclose(float(actual), expected, rel_tol=tolerance, abs_tol=tolerance):
        raise AssertionError(f"{label}: {actual} != {expected}")


def read_sheet(wb, name):
    rows = wb[name].iter_rows(values_only=True)
    columns = next(rows)
    return [dict(zip(columns, row)) for row in rows]


def main():
    source = ROOT / "data" / "理论-实验对比结果.xlsx"
    w = load_workbook(source, read_only=True, data_only=True)
    c = read_sheet(w, "分档曲线_可直接作图")
    p = read_sheet(w, "逐点明细")
    w.close()
    detail = read_csv("condition_details.csv")
    table = read_csv("table3_revised_unrounded.csv")
    old = read_csv("table3_legacy_reproduced.csv")
    assert len(c)==len(detail)==231
    assert len(p)==877
    assert sum(int(x["n"]) for x in c)==len(p)
    assert len({(r['rpm_nom'],r['alpha_deg'],r['J']) for r in c})==231
    for i,(src,d) in enumerate(zip(c,detail)):
        assert int(d['source_curve_excel_row'])==i+2
        for field in ['rpm_nom','alpha_deg','J','C_T_exp','C_T_theory','n']:
            close(d[field],float(src[field]),f"row {i+2} frozen {field}")
        measured,predicted=src['C_T_exp'],src['C_T_theory']
        assert math.isfinite(measured) and measured>0
        close(d['residual_CT'],measured-predicted,f"row {i+2} residual")
        close(d['signed_relative_error_pct'],100*(measured-predicted)/measured,f"row {i+2} relative")
    for i,a in enumerate([0,15,30,45,60,75,90,None]):
        group = [r for r in c if a is None or r['alpha_deg']==a]
        assert len(group)==(231 if a is None else 33)
        mean=statistics.fmean(r['C_T_exp'] for r in group)
        deltas=[r['C_T_exp']-r['C_T_theory'] for r in group]
        rmse=math.sqrt(math.fsum(x*x for x in deltas)/len(group))
        rel=[100*x/r['C_T_exp'] for x,r in zip(deltas,group)]
        row=table[i]
        close(row['mean_CT_exp'],mean,f"{a} mean")
        close(row['RMSE_CT'],rmse,f"{a} RMSE")
        close(row['NRMSE_pct'],100*rmse/mean,f"{a} NRMSE")
        close(row['normalized_max_abs_error_pct'],100*max(map(abs,deltas))/mean,f"{a} max")
        close(row['MAPE_pct'],statistics.fmean(map(abs,rel)),f"{a} MAPE")
        close(row['median_signed_relative_error_pct'],statistics.median(rel),f"{a} median")
        raw=[r for r in p if a is None or r['alpha_deg']==a]
        usable=[r for r in raw if r['thrust_ok'] is True]
        close(old[i]['rmse_norm_pct'],math.sqrt(statistics.fmean(r['err_norm_pct']**2 for r in raw)),f"{a} old RMSE")
        close(old[i]['maxerr_norm_pct'],max(abs(r['err_norm_pct']) for r in raw),f"{a} old max")
        close(old[i]['mape_pct'],statistics.fmean(abs(r['err_rel_pct']) for r in usable),f"{a} old MAPE")
    pooled_nrmse=float(table[-1]['NRMSE_pct'])
    assert not math.isclose(pooled_nrmse,statistics.fmean(float(r['NRMSE_pct']) for r in table[:-1]),abs_tol=.01)
    rel=[abs(100*(r['C_T_exp']-r['C_T_theory'])/r['C_T_exp']) for r in c]
    assert sum(v<=5 for v in rel)==144
    assert sum(v<=10 for v in rel)==197
    assert sum(r['thrust_ok'] is True for r in p)==869
    assert [f"{float(old[-1][field]):.2f}" for field in ['rmse_norm_pct','mape_pct','maxerr_norm_pct']]==['5.76','5.08','19.75']
    # Each parity-boundary construction must give exactly its intended measured-relative error.
    for fraction in [.05,.10]:
        for sign in [-1,1]:
            for x in [.002,.008,.015]:
                y=x/(1+sign*fraction)
                close(abs(y-x)/y,fraction,"parity guide measured denominator")
    protected=json.loads((OUT/'protected_inputs_sha256.json').read_text(encoding='utf-8'))
    baseline=json.loads((OUT/'before_20261013'/'manifest.json').read_text(encoding='utf-8'))
    for name,digest in protected.items():
        actual=hashlib.sha256((ROOT/name).read_bytes()).hexdigest()
        assert actual==digest==baseline[name.replace('\\','/')],f"Protected input changed: {name}"
    for name,digest in baseline.items():
        assert hashlib.sha256((OUT/'before_20261013'/name).read_bytes()).hexdigest()==digest,f"Archive changed: {name}"
    # Independently isolate only the two explicitly authorized manuscript blocks.
    def outside_blocks(data):
        spans=[]
        for marker in [b'{Image/fig9_rotor_model_diagnostics.pdf}', '{\\small 表3\\quad'.encode('utf-8')]:
            assert data.count(marker)==1
            i=data.index(marker)
            start=data.rfind(b'\\begin{center}',0,i)
            end=data.index(b'\\end{center}',i)+len(b'\\end{center}')
            spans.append((start,end))
        spans.sort()
        return data[:spans[0][0]]+data[spans[0][1]:spans[1][0]]+data[spans[1][1]:]
    current_tex=(ROOT/'template.tex').read_bytes()
    archived_tex=(OUT/'before_20261013/template.tex').read_bytes()
    assert outside_blocks(current_tex)==outside_blocks(archived_tex),'Manuscript changed outside authorized blocks'
    # Confirm source and canonical vector outputs match, with no embedded bitmap marks.
    import fitz
    doc=fitz.open(OUT/'fig7_revised.pdf')
    assert len(doc)==1 and not doc[0].get_images()
    text=doc[0].get_text()
    assert 'within' not in text and '85.3%' not in text and '10%' in text
    assert all(word not in text for word in ['CCC','Pearson','R²'])
    assert len(doc[0].get_drawings())>100
    doc.close()
    for suffix in ['pdf','png']:
        assert (OUT/f'fig7_revised.{suffix}').read_bytes()==(ROOT/'Image'/f'fig9_rotor_model_diagnostics.{suffix}').read_bytes()
    assert (OUT/'table3_revised.tex').read_bytes()==(ROOT/'data_process'/'model_validation_table.tex').read_bytes()
    checks={"status":"PASS", "independent_source_recalculation":"all 231 CT pairs; all eight new/legacy table rows",
            "threshold_counts":{"5_percent":144,"10_percent":197,"N":231},
            "protected_inputs_and_archive":"SHA256 unchanged", "manuscript":"outside Figure 7 caption / Table 3 byte-identical",
            "vector_pdf":"one page, no embedded images",
            "limitation":"This validates frozen-pair arithmetic, not upstream density/rpm/noise consistency."}
    (OUT/'verification.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(checks,ensure_ascii=False,indent=2))


if __name__=='__main__':
    main()
